import json
import os
import openpyxl
import re
from collections import Counter

BASE_DIR = r"d:\RetroVault_App"
with open(os.path.join(BASE_DIR, '_internal', 'snes_libretro_tree.json'), 'r', encoding='utf-8') as f:
    tree = json.load(f)

boxarts_by_name = {
    os.path.basename(item['path'])[:-4].lower(): item['path']
    for item in tree.get('tree', [])
    if item['path'].startswith('Named_Boxarts/')
}

wb = openpyxl.load_workbook(os.path.join(BASE_DIR, 'data', '收藏.xlsx'), data_only=True)
ws = wb['SFC']
games = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(vals) and vals[0] is not None:
        games.append({
            'no': vals[0],
            'jp': str(vals[1] or '').strip(),
            'zh': str(vals[2] or '').strip(),
            'en': str(vals[3] or '').strip(),
            'code': str(vals[4] or '').strip(),
            'region': str(vals[5] or '').strip(),
            'genre_raw': str(vals[6] or '').strip(),
        })

EXPLICIT_MAPPING = {
    'SHVC-TK-1': 'Teitoku no Ketsudan (Japan)',
    'SHVC-8X': 'Super Donkey Kong (Japan)',
    'SHVC-AKFE': 'Hoshi no Kirby Super Deluxe (Japan)',
    'SHVC-P2': 'Jissen Pachinko Hisshouhou! 2 (Japan)',
    'SHVC-4M': 'Super Mario Collection (Japan)',
    'SHVC-EN': 'Elnard (Japan)',
    'SHVC-K2': 'Seiken Densetsu 2 (Japan)',
    'SHVC-A3TJ-JPN': 'Seiken Densetsu 3 (Japan)',
    'SHVC-AQ8J-JPN': 'Tenchi Souzou (Japan)',
    'SHVC-AL': 'Air Management - Oozora ni Kakeru (Japan)',
    'SHVC-GA': 'New 3D Golf Simulation - Harukanaru Augusta (Japan)',
    'SHVC-O2': 'Masters - Harukanaru Augusta 2 (Japan)',
    'SHVC-IZ': 'Super Nobunaga no Yabou - Bushou Fuuunroku (Japan)',
    'SHVC-NK': 'Nobunaga no Yabou - Haouden (Japan)',
    'SHVC-A98J': 'Tokimeki Memorial - Densetsu no Ki no Shita de (Japan)',
    'SHVC-AH3J-JPN': 'Nobunaga no Yabou - Tenshouki (Japan)',
    'SHVC-2Z': 'Dragon Ball Z - Super Butouden (Japan)',
    'SHVC-DB': 'Dragon Ball Z - Super Saiya Densetsu (Japan)',
    'SHVC-JE': 'J.League Soccer - Prime Goal (Japan)',
    'SHVC-OJ': 'Garou Densetsu 2 - Aratanaru Tatakai (Japan)',
    'SHVC-N3': 'Feda - The Emblem of Justice (Japan)',
    'SHVC-A4GJ-JPN': 'Ganbare Goemon Kirakira Douchuu - Boku ga Dancer ni Natta Wake (Japan)',
    'SHVC-S3': 'Sangokushi III (Japan)',
    'SHVC-8L-1': 'Super Sangokushi II (Japan)',
    'SHVC-I4': 'Super Sangokushi (Japan)',
    'SHVC-DQ': 'Dragon Quest I _ II (Japan)',
    'SHVC-E7': 'Dragon Ball Z - Super Butouden 2 (Japan)',
    'SHVC-3M': 'Dragon Ball Z - Super Butouden 3 (Japan)',
    'SHVC-2N': 'J.League Soccer - Prime Goal 2 (Japan)',
    'SHVC-AJ3J-JPN': 'J.League Soccer - Prime Goal 3 (Japan)',
    'SHVC-Y1': 'Yu Yu Hakusho - Tokubetsu Hen (Japan)',
    'SHVC-Y2': 'Yu Yu Hakusho 2 - Kakutou no Shou (Japan)',
    'SHVC-AYYJ-JPN': 'Yu Yu Hakusho Final - Makai Saikyou Retsuden (Japan)',
    'SHVC-ADNJ': 'Super Donkey Kong 2 - Dixie _ Diddy (Japan)',
    'SHVC-A3CJ': 'Super Donkey Kong 3 - Nazo no Kremis-tou (Japan)',
    'SHVC-FN': 'F-1 Grand Prix - Part II (Japan)',
    'SHVC-A2SJ-JPN': 'Sword World SFC 2 - Inishie no Kyojin Densetsu (Japan)',
    'SHVC-ANIJ-JPN': 'Estpolis Denki II (Japan)',
    'SHVC-AFGJ-JPN': 'Front Mission Series - Gun Hazard (Japan)',
    'SHVC-P-ASEJ': 'Yokoyama Mitsuteru Sangokushi Bangi - Sugoroku Eiyuu-ki (Japan)',
    'SHVC-G8': 'Daibakushou Jinsei Gekijou - Dokidoki Seishun Hen (Japan)',
    'SHVC-MW': 'Super Mario World - Super Mario Bros. 4 (Japan)',
    'SHVC-P-APXJ': 'Parlor! Parlor! (Japan)',
    'SHVC-B9': 'Drift King Shutokou Battle 2 - Tsuchiya Keiichi _ Bandou Masaaki (Japan)',
    'SHVC-W3': 'Sangokushi Seishi - Tenbu Spirits (Japan)',
    'SHVC-S8': 'Silva Saga II - The Legend of Light and Darkness (Japan)',
    'SHVC-OB': 'Densetsu no Ogre Battle - The March of the Black Queen (Japan)',
    'SHVC-QI': 'Dokuritsu Sensou - Liberty or Death (Japan)',
    'SHVC-PM': 'PowerMonger - Mashou no Bouryaku (Japan)',
    'SHVC-8K': 'Garou Densetsu - Shukumei no Tatakai (Japan)',
    'SHVC-28': 'Garou Densetsu Special (Japan)',
    'SHVC-JQ': 'Derby Jockey - Kishu Ou e no Michi (Japan)',
    'SHVC-P-AMDJ': 'Super Momotarou Dentetsu III (Japan)',
    'SHVC-P-ATXJ': 'Super Momotarou Dentetsu DX (Japan)',
    'SHVC-DS': 'Daisenryaku Expert (Japan)',
    'SHVC-IK': 'Super Honmei - GI Seiha (Japan)',
    'SHVC-UG': 'Yumemaboroshi no Gotoku (Japan)',
    'SHVC-D3': 'Derby Stallion III (Japan)',
    'SHVC-JA': 'Super Famista (Japan)',
    'SHVC-FI': 'Super Famista 2 (Japan)',
    'SHVC-FS': 'Honke Sankyo Fever - Jikki Simulation (Japan)',
    'SHVC-GC': 'New 3D Golf Simulation - Waialae no Kiseki (Japan)',
    'SHVC-PQ': 'Super Puyo Puyo (Japan)',
    'SHVC-T2': 'Super Tetris 2 + Bombliss (Japan) (En)',
    'SHVC-IQ': 'Super Mahjong Taikai (Japan)',
    'SHVC-ARVJ': 'Tales of Phantasia (Japan)',
    'SHVC-WJ': 'Winning Post (Japan)',
    'SHVC-P-AWPJ': 'Winning Post 2 (Japan)',
    'SHVC-P-AOBJ': 'Stable Star - Kyuusha Monogatari (Japan)',
    'SHVC-B2': 'Thoroughbred Breeder II (Japan)',
    'SHVC-E4': 'Super Hanafuda (Japan)',
    'SHVC-TR': 'Torneko no Daibouken - Fushigi no Dungeon (Japan)',
    'SHVC-D6': 'Dragon Quest VI - Maboroshi no Daichi (Japan)',
    'SHVC-D5': 'Dragon Quest V - Tenkuu no Hanayome (Japan)',
    'SHVC-F5': 'Final Fantasy V (Japan)',
    'SHVC-F6': 'Final Fantasy VI (Japan)',
    'SHVC-4F': 'Final Fantasy IV (Japan)',
    'SHVC-CT': 'Chrono Trigger (Japan)',
    'SHVC-SE': 'Seiken Densetsu 2 (Japan)',
    'SHVC-R2': 'Romancing Sa-Ga 2 (Japan)',
    'SHVC-R3': 'Romancing Sa-Ga 3 (Japan)',
    'SHVC-RS': 'Romancing Sa-Ga (Japan)',
    'SHVC-MK': 'Super Mario Kart (Japan)',
    'SHVC-2M': 'Super Mario Collection (Japan)',
    'SHVC-YI': 'Super Mario - Yoshi Island (Japan)',
    'SHVC-P-AKBJ': 'Konpeki no Kantai (Japan)',
    'SHVC-AR3J': 'Romancing Sa-Ga 3 (Japan)',
    'SHVC-NZ': 'Super Nobunaga no Yabou - Zenkoku Ban (Japan)',
    'SHVC-LN': 'Lennus - Kodai Kikai no Kioku (Japan)',
    'SHVC-SK': 'Super Aoki Ookami to Shiroki Mejika - Genchou Hishi (Japan)',
    'SHVC-RT': 'RPG Tsukuru - Super Dante (Japan)',
    'SHVC-AZ8J-JPN': 'Dragon Ball Z - Super Gokuu Den - Totsugeki Hen (Japan)',
    'SHVC-AK8J-JPN': 'Dragon Ball Z - Super Gokuu Den - Kakusei Hen (Japan)',
    'SHVC-ND': 'Kunio-kun no Dodge Ball Da yo Zenin Shuugou! (Japan)',
    'SHVC-5K': 'Ranma 1-2 - Chougi Ranbu Hen (Japan)',
    'SHVC-64': 'Ranma 1-2 - Akanekodan Teki Hihou (Japan)',
    'SHVC-RM': 'Ranma 1-2 - Chounai Gekitou Hen (Japan)',
    'SHVC-RJ': 'Ranma 1-2 - Bakuretsu Rantou Hen (Japan)',
    'SHVC-ED': 'Dragon Slayer - Eiyuu Densetsu (Japan)',
    'SHVC-HA': 'Super Mahjong 2 - Honkaku 4-nin Uchi! (Japan)',
    'SHVC-PW': 'Shin Nihon Pro Wrestling - Chou Senshi in Tokyo Dome - Fantastic Story (Japan)',
    'SHVC-ZL': 'Zelda no Densetsu - Kamigami no Triforce (Japan)',
    'SHVC-W5': 'Wizardry V - Saika no Chuushin (Japan) (En,Ja)',
    'SHVC-A4KJ-JPN': 'Mini Yonku Shining Scorpion - Let\'s _ Go!! (Japan)',
    'SHVC-U2': 'Ultra Baseball Jitsumei Ban (Japan)',
    'SHVC-HI': 'Super Kyuukyoku Harikiri Stadium (Japan)',
    'SHVC-3G': 'Ganbare Goemon 3 - Shishi Juurokubee no Karakuri Manjigatame (Japan)',
    'SHVC-HP': 'Hiryuu no Ken S - Hyper Version (Japan)',
    'SHVC-GG': 'Ganbare Goemon - Yuki Hime Kyuushutsu Emaki (Japan)'
}

def clean_key(s):
    s = s.replace('&', 'and').replace('ou', 'o').replace('uu', 'u')
    s = re.sub(r'[\(\)\[\]（）「」『』～\-:：!！\?\？\+＆_·\.\',]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip().lower()

matches = {}
for g in games:
    code = g['code']
    en = g['en']
    jp = g['jp']
    found = None
    if code in EXPLICIT_MAPPING:
        target = EXPLICIT_MAPPING[code].lower()
        if target in boxarts_by_name:
            found = boxarts_by_name[target]
    
    if not found:
        for c in [f"{en} (Japan)", f"{en} (Japan) (En)", f"{en} (USA)", f"{jp} (Japan)", en]:
            if c.lower() in boxarts_by_name:
                found = boxarts_by_name[c.lower()]
                break
                
    if not found and en:
        c_en = clean_key(en)
        for b_name, b_path in boxarts_by_name.items():
            if '(japan)' in b_name:
                b_clean = clean_key(b_name.replace('(japan)', ''))
                if c_en == b_clean:
                    found = b_path
                    break
    matches[code] = found

missing = [g for g in games if not matches.get(g['code'])]
print(f"Total games: {len(games)}, Matched: {len(games) - len(missing)}, Missing: {len(missing)}")
if missing:
    for m in missing:
        print(f"  MISSING: [{m['code']}] {m['jp']} | {m['zh']} | {m['en']}")

counts = Counter(p for p in matches.values() if p)
dups = {p: cnt for p, cnt in counts.items() if cnt > 1}
print(f"\nDuplicates count: {len(dups)}")
for p, cnt in dups.items():
    gms = [g for g in games if matches.get(g['code']) == p]
    print(f"  DUP ({cnt}x): {p}")
    for g in gms:
        print(f"    [{g['code']}] {g['jp']}")
