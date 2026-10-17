import os
import sys
import re
import csv
import json
import time
import uuid
import urllib.request
import urllib.parse
from difflib import SequenceMatcher
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    import openpyxl
except ImportError:
    openpyxl = None

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DATA_DIR = os.path.join(BASE_DIR, 'data')
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
CDN_DIR = os.path.join(DOCS_DIR, 'cdn')
INTERNAL_DIR = os.path.join(BASE_DIR, '_internal')

os.makedirs(CDN_DIR, exist_ok=True)

# 1. Load Libretro SNES Tree
with open(os.path.join(INTERNAL_DIR, 'snes_libretro_tree.json'), 'r', encoding='utf-8') as f:
    tree_data = json.load(f)

boxarts_by_name = {}
snaps_by_name = {}
for item in tree_data.get('tree', []):
    p = item['path']
    if p.startswith('Named_Boxarts/'):
        name = os.path.basename(p)[:-4]
        boxarts_by_name[name.lower()] = p
    elif p.startswith('Named_Snaps/'):
        name = os.path.basename(p)[:-4]
        snaps_by_name[name.lower()] = p

# Explicit mappings for manual matching
EXPLICIT_MAPPING = {
    # Augusta Golf series - distinct covers
    'SHVC-GA': 'New 3D Golf Simulation - Harukanaru Augusta (Japan)',
    'SHVC-O2': 'Masters - Harukanaru Augusta 2 (Japan)',
    # Previously 404 / broken covers replaced with verified Japan scans
    'SHVC-IZ': 'Super Nobunaga no Yabou - Bushou Fuuunroku (Japan)',
    'SHVC-NK': 'Nobunaga no Yabou - Haouden (Japan)',
    'SHVC-A98J': 'Tokimeki Memorial - Densetsu no Ki no Shita de (Japan)',
    'SHVC-AH3J-JPN': 'Nobunaga no Yabou - Tenshouki (Japan)',
    'SHVC-2Z': 'Dragon Ball Z - Super Butouden (Japan)',
    'SHVC-DB': 'Dragon Ball Z - Super Saiya Densetsu (Japan)',
    'SHVC-JE': 'J.League Soccer - Prime Goal (Japan)',
    'SHVC-OJ': 'Garou Densetsu 2 - Aratanaru Tatakai (Japan)',
    'SHVC-N3': 'Feda - The Emblem of Justice (Japan)',
    'SHVC-A4GJ-JPN': 'Ganbare Goemon Kirakira Douchuu - Boku ga Dancer ni Natta Wake (Japan)',
    # Series disambiguation (fixing sequel-prequel collision)
    'SHVC-S3': 'Sangokushi III (Japan)',
    'SHVC-8L-1': 'Super Sangokushi II (Japan)',
    'SHVC-I4': 'Super Sangokushi (Japan)',
    'SHVC-DQ': 'Dragon Quest I _ II (Japan)',
    'SHVC-D5': 'Dragon Quest V - Tenkuu no Hanayome (Japan)',
    'SHVC-D6': 'Dragon Quest VI - Maboroshi no Daichi (Japan)',
    'SHVC-E7': 'Dragon Ball Z - Super Butouden 2 (Japan)',
    'SHVC-3M': 'Dragon Ball Z - Super Butouden 3 (Japan)',
    'SHVC-2N': 'J.League Soccer - Prime Goal 2 (Japan)',
    'SHVC-AJ3J-JPN': 'J.League Soccer - Prime Goal 3 (Japan)',
    'SHVC-Y1': 'Yu Yu Hakusho - Tokubetsu Hen (Japan)',
    'SHVC-Y2': 'Yu Yu Hakusho 2 - Kakutou no Shou (Japan)',
    'SHVC-AYYJ-JPN': 'Yu Yu Hakusho Final - Makai Saikyou Retsuden (Japan)',
    'SHVC-8X': 'Super Donkey Kong (Japan)',
    'SHVC-ADNJ': 'Super Donkey Kong 2 - Dixie _ Diddy (Japan)',
    'SHVC-A3CJ': 'Super Donkey Kong 3 - Nazo no Kremis-tou (Japan)',
    'SHVC-FN': 'F-1 Grand Prix - Part II (Japan)',
    'SHVC-A2SJ-JPN': 'Sword World SFC 2 - Inishie no Kyojin Densetsu (Japan)',
    'SHVC-ANIJ-JPN': 'Estpolis Denki II (Japan)',
    'SHVC-AFGJ-JPN': 'Front Mission Series - Gun Hazard (Japan)',
    'SHVC-P-ASEJ': 'Yokoyama Mitsuteru Sangokushi Bangi - Sugoroku Eiyuu-ki (Japan)',
    'SHVC-G8': 'Daibakushou Jinsei Gekijou - Dokidoki Seishun Hen (Japan)',
    'SHVC-MW': 'Super Mario World - Super Mario Bros. 4 (Japan)',
    'SHVC-P-APXJ': 'Parlor! Parlor! (Japan)',
    'SHVC-B9': 'Drift King Shutokou Battle 2 - Tsuchiya Keiichi _ Bandou Masaaki (Japan)',
    'SHVC-W3': 'Sangokushi Seishi - Tenbu Spirits (Japan)',
    'SHVC-S8': 'Silva Saga II - The Legend of Light and Darkness (Japan)',
    'SHVC-OB': 'Densetsu no Ogre Battle - The March of the Black Queen (Japan)',
    'SHVC-QI': 'Dokuritsu Sensou - Liberty or Death (Japan)',
    'SHVC-PM': 'PowerMonger - Mashou no Bouryaku (Japan)',
    'SHVC-8K': 'Garou Densetsu - Shukumei no Tatakai (Japan)',
    'SHVC-28': 'Garou Densetsu Special (Japan)',
    'SHVC-TK': 'Teitoku no Ketsudan (Japan)',
    'SHVC-TK-1': 'Teitoku no Ketsudan (Japan)',
    'SHVC-AKFE': 'Hoshi no Kirby Super Deluxe (Japan)',
    'SHVC-P2': 'Jissen Pachinko Hisshouhou! 2 (Japan)',
    'SHVC-4M': 'Super Mario Collection (Japan)',
    'SHVC-EN': 'Elnard (Japan)',
    'SHVC-K2': 'Seiken Densetsu 2 (Japan)',
    'SHVC-A3TJ-JPN': 'Seiken Densetsu 3 (Japan)',
    'SHVC-AQ8J-JPN': 'Tenchi Souzou (Japan)',
    'SHVC-AL': 'Air Management - Oozora ni Kakeru (Japan)',
    'SHVC-JQ': 'Derby Jockey - Kishu Ou e no Michi (Japan)',
    'SHVC-P-AMDJ': 'Super Momotarou Dentetsu III (Japan)',
    'SHVC-P-ATXJ': 'Super Momotarou Dentetsu DX (Japan)',
    'SHVC-DS': 'Daisenryaku Expert (Japan)',
    'SHVC-IK': 'Super Honmei - GI Seiha (Japan)',
    'SHVC-UG': 'Yumemaboroshi no Gotoku (Japan)',
    'SHVC-D3': 'Derby Stallion III (Japan)',
    'SHVC-JA': 'Super Famista (Japan)',
    'SHVC-FI': 'Super Famista 2 (Japan)',
    'SHVC-FS': 'Honke Sankyo Fever - Jikki Simulation (Japan)',
    'SHVC-GC': 'New 3D Golf Simulation - Waialae no Kiseki (Japan)',
    'SHVC-PQ': 'Super Puyo Puyo (Japan)',
    'SHVC-T2': 'Super Tetris 2 + Bombliss (Japan) (En)',
    'SHVC-IQ': 'Super Mahjong Taikai (Japan)',
    'SHVC-ARVJ': 'Tales of Phantasia (Japan)',
    'SHVC-WJ': 'Winning Post (Japan)',
    'SHVC-P-AWPJ': 'Winning Post 2 (Japan)',
    'SHVC-P-AOBJ': 'Stable Star - Kyuusha Monogatari (Japan)',
    'SHVC-B2': 'Thoroughbred Breeder II (Japan)',
    'SHVC-E4': 'Super Hanafuda (Japan)',
    'SHVC-TR': 'Torneko no Daibouken - Fushigi no Dungeon (Japan)',
    'SHVC-F5': 'Final Fantasy V (Japan)',
    'SHVC-F6': 'Final Fantasy VI (Japan)',
    'SHVC-4F': 'Final Fantasy IV (Japan)',
    'SHVC-CT': 'Chrono Trigger (Japan)',
    'SHVC-SE': 'Seiken Densetsu 2 (Japan)',
    'SHVC-R2': 'Romancing Sa-Ga 2 (Japan)',
    'SHVC-R3': 'Romancing Sa-Ga 3 (Japan)',
    'SHVC-RS': 'Romancing Sa-Ga (Japan)',
    'SHVC-MK': 'Super Mario Kart (Japan)',
    'SHVC-2M': 'Super Mario Collection (Japan)',
    'SHVC-YI': 'Super Mario - Yoshi Island (Japan)',
    'SHVC-P-AKBJ': 'Konpeki no Kantai (Japan)',
    'SHVC-AR3J': 'Romancing Sa-Ga 3 (Japan)',
    'SHVC-NZ': 'Super Nobunaga no Yabou - Zenkoku Ban (Japan)',
    'SHVC-LN': 'Lennus - Kodai Kikai no Kioku (Japan)',
    'SHVC-SK': 'Super Aoki Ookami to Shiroki Mejika - Genchou Hishi (Japan)',
    'SHVC-RT': 'RPG Tsukuru - Super Dante (Japan)',
    'SHVC-AZ8J-JPN': 'Dragon Ball Z - Super Gokuu Den - Totsugeki Hen (Japan)',
    'SHVC-AK8J-JPN': 'Dragon Ball Z - Super Gokuu Den - Kakusei Hen (Japan)',
    'SHVC-ND': 'Kunio-kun no Dodge Ball Da yo Zenin Shuugou! (Japan)',
    'SHVC-5K': 'Ranma 1-2 - Chougi Ranbu Hen (Japan)',
    'SHVC-64': 'Ranma 1-2 - Akanekodan Teki Hihou (Japan)',
    'SHVC-RM': 'Ranma 1-2 - Chounai Gekitou Hen (Japan)',
    'SHVC-RJ': 'Ranma 1-2 - Bakuretsu Rantou Hen (Japan)',
    'SHVC-ED': 'Dragon Slayer - Eiyuu Densetsu (Japan)',
    'SHVC-HA': 'Super Mahjong 2 - Honkaku 4-nin Uchi! (Japan)',
    'SHVC-PW': 'Shin Nihon Pro Wrestling - Chou Senshi in Tokyo Dome - Fantastic Story (Japan)',
    'SHVC-ZL': 'Zelda no Densetsu - Kamigami no Triforce (Japan)',
    'SHVC-W5': 'Wizardry V - Saika no Chuushin (Japan) (En,Ja)',
    'SHVC-A4KJ-JPN': 'Mini Yonku Shining Scorpion - Let\'s _ Go!! (Japan)',
    'SHVC-U2': 'Ultra Baseball Jitsumei Ban (Japan)',
    'SHVC-HI': 'Super Kyuukyoku Harikiri Stadium (Japan)',
    'SHVC-3G': 'Ganbare Goemon 3 - Shishi Juurokubee no Karakuri Manjigatame (Japan)',
    'SHVC-HP': 'Hiryuu no Ken S - Hyper Version (Japan)',
    'SHVC-GG': 'Ganbare Goemon - Yuki Hime Kyuushutsu Emaki (Japan)',
}

def clean_key(s):
    s = s.replace('&', 'and').replace('ou', 'o').replace('uu', 'u')
    s = re.sub(r'[\(\)\[\]（）「」『』～\-:：!！\?\？\+＆_·\.\',]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip().lower()

def match_libretro(game):
    code = game['code']
    en = game['en']
    jp = game['jp']

    found_box = None
    if code in EXPLICIT_MAPPING:
        target = EXPLICIT_MAPPING[code].lower()
        if target in boxarts_by_name:
            found_box = boxarts_by_name[target]

    if not found_box:
        for c in [f"{en} (Japan)", f"{en} (Japan) (En)", f"{en} (USA)", f"{jp} (Japan)", en]:
            if c.lower() in boxarts_by_name:
                found_box = boxarts_by_name[c.lower()]
                break

    if not found_box and en:
        c_en = clean_key(en)
        for b_name, b_path in boxarts_by_name.items():
            if '(japan)' in b_name:
                b_clean = clean_key(b_name.replace('(japan)', ''))
                if c_en == b_clean:
                    found_box = b_path
                    break

    found_snap = None
    if found_box:
        base = os.path.basename(found_box)[:-4].lower()
        if base in snaps_by_name:
            found_snap = snaps_by_name[base]

    return found_box, found_snap

# 2. LaunchBox Scraper Helpers
HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
}

def clean_lb_query(text):
    if not text:
        return ""
    t = re.split(r'[:～\-\(（]', text)[0].strip()
    t = re.sub(r'[\'\"・\.]', ' ', t)
    return ' '.join(t.split())

def query_from_boxart(boxart_path):
    if not boxart_path:
        return ""
    fn = os.path.basename(boxart_path)[:-4]
    fn = re.sub(r'\(.*?\)', '', fn).strip()
    return fn

def similarity(a, b):
    return SequenceMatcher(None, a.lower(), b.lower()).ratio()

def search_launchbox(query, platform_name='Super Nintendo Entertainment System'):
    if not query:
        return []
    enc_q = urllib.parse.quote(query)
    enc_plat = urllib.parse.quote(platform_name)
    url = f"https://gamesdb.launchbox-app.com/games/results/{enc_q}?platform={enc_plat}&searchText={enc_q}"
    req = urllib.request.Request(url, headers=HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=12) as resp:
            html = resp.read().decode('utf-8', errors='ignore')
    except Exception:
        return []

    pattern = re.compile(
        r'<a[^>]+href=["\'](/games/details/\d+-[^"\']*)["\'][^>]*>.*?'
        r'<div[^>]*class="cardTitle"[^>]*>\s*<h3[^>]*>(.*?)</h3>\s*<p[^>]*>(.*?)</p>',
        re.DOTALL
    )
    results = []
    for m in pattern.finditer(html):
        link = m.group(1)
        title = m.group(2).strip()
        plat = m.group(3).strip()
        if 'super nintendo' in plat.lower() or 'snes' in plat.lower():
            results.append({
                'link': link,
                'title': title,
                'platform': plat
            })
    return results

def get_launchbox_images(detail_url_path):
    url = f"https://gamesdb.launchbox-app.com{detail_url_path}"
    req = urllib.request.Request(url, headers=HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=12) as resp:
            html = resp.read().decode('utf-8', errors='ignore')
    except Exception:
        return {}

    img_pattern = re.compile(r'<img[^>]+src=["\'](https://images\.launchbox-app\.com/[^"\']+)["\'][^>]*alt=["\']([^"\']*)["\']')
    images = []
    for m in img_pattern.finditer(html):
        src = m.group(1).replace('.com//', '.com/')
        alt = m.group(2)
        images.append({'src': src, 'alt': alt})

    def pick_best(img_type):
        candidates = [img for img in images if img_type.lower() in img['alt'].lower()]
        if not candidates:
            return None
        region_order = ['(japan)', '(asia)', '(world)', '(north america)', '(united states)', '(europe)']
        for reg in region_order:
            for c in candidates:
                if reg in c['alt'].lower():
                    return c
        return candidates[0]

    return {
        'back': pick_best('Box - Back'),
        'spine': pick_best('Box - Spine'),
        'front': pick_best('Box - Front')
    }

def download_cdn_image(src_url):
    """Download image to docs/cdn/<filename> if not present"""
    if not src_url:
        return None
    fn = os.path.basename(urllib.parse.urlparse(src_url).path)
    dest = os.path.join(CDN_DIR, fn)
    if os.path.exists(dest) and os.path.getsize(dest) > 500:
        return f"https://raw.githubusercontent.com/Kevinte67228/Retro-library/main/docs/cdn/{fn}"

    req = urllib.request.Request(src_url, headers=HEADERS)
    for _ in range(3):
        try:
            with urllib.request.urlopen(req, timeout=15) as resp:
                if resp.status == 200:
                    data = resp.read()
                    if len(data) > 0:
                        with open(dest, 'wb') as f:
                            f.write(data)
                        return f"https://raw.githubusercontent.com/Kevinte67228/Retro-library/main/docs/cdn/{fn}"
        except Exception:
            time.sleep(0.5)
    return None

# Metadata Parsing Helpers
def extract_developer_publisher(notes):
    m = re.search(r'\((.*?)\)', notes)
    inner = m.group(1).strip() if m else notes.strip()
    
    companies = [
        ('Nintendo', '任天堂 (Nintendo)', '任天堂 (Nintendo)'),
        ('任天堂', '任天堂 (Nintendo)', '任天堂 (Nintendo)'),
        ('Square', 'スクウェア (Square)', 'スクウェア (Square)'),
        ('Capcom', 'カプコン (Capcom)', 'カプコン (Capcom)'),
        ('Enix', 'エニックス (Enix)', 'エニックス (Enix)'),
        ('Koei', '光栄 (Koei)', '光栄 (Koei)'),
        ('光榮', '光栄 (Koei)', '光栄 (Koei)'),
        ('Konami', 'コナミ (Konami)', 'コナミ (Konami)'),
        ('Bandai', 'バンダイ (Bandai)', 'バンダイ (Bandai)'),
        ('Hudson', 'ハドソン (Hudson)', 'ハドソン (Hudson)'),
        ('Namco', 'ナムコ (Namco)', 'ナムコット (Namcot)'),
        ('Namcot', 'ナムコ (Namco)', 'ナムコット (Namcot)'),
        ('Taito', 'タイトー (Taito)', 'タイトー (Taito)'),
        ('Banpresto', 'バンプレスト (Banpresto)', 'バンプレスト (Banpresto)'),
        ("I'MAX", "アイマックス (I'MAX)", "アイマックス (I'MAX)"),
        ('ASCII', 'アスキー (ASCII)', 'アスキー (ASCII)'),
        ('T&E Soft', 'ティーアンドイーソフト (T&E Soft)', 'ティーアンドイーソフト (T&E Soft)'),
        ('Jaleco', 'ジャレコ (Jaleco)', 'ジャレコ (Jaleco)'),
        ('Human', 'ヒューマン (Human)', 'ヒューマン (Human)'),
        ('Data East', 'データイースト (Data East)', 'データイースト (Data East)'),
        ('Telenet', '日本テレネット (Telenet Japan)', '日本テレネット (Telenet Japan)'),
        ('Chunsoft', 'チュンソフト (Chunsoft)', 'チュンソフト (Chunsoft)'),
        ('Wolf Team', 'ウルフ・チーム (Wolf Team)', '日本テレネット (Telenet Japan)'),
        ('Athena', 'アテナ (Athena)', 'アテナ (Athena)'),
        ('Angel', 'エンジェル (Angel)', 'エンジェル (Angel)'),
        ('Yutaka', 'ユタカ (Yutaka)', 'ユタカ (Yutaka)'),
        ('Epoch', 'エポック社 (Epoch)', 'エポック社 (Epoch)'),
    ]
    
    for kw, dev, pub in companies:
        if kw.lower() in inner.lower():
            if '/' in inner:
                parts = [p.strip() for p in inner.split('/')]
                return f"{parts[0]}", pub
            return dev, pub
    
    return inner if inner else 'Nintendo', inner if inner else 'Nintendo'

def extract_genre(notes):
    n = notes.lower()
    if '動作rpg' in n or '動作角色扮演' in n or 'arpg' in n:
        return 'ARPG 動作角色扮演'
    if 'rpg' in n or '角色扮演' in n:
        return 'RPG 角色扮演'
    if 'act' in n or '動作' in n or '橫向' in n or '過關' in n:
        return 'ACT 動作'
    if 'ftg' in n or '格鬥' in n or '對戰' in n:
        return 'FTG 格鬥'
    if 'slg' in n or '模擬' in n or '策略' in n or '戰略' in n or '戰棋' in n or '歷史' in n:
        return 'SLG 模擬'
    if 'stg' in n or '射擊' in n or '飛行' in n:
        return 'STG 射擊'
    if '賽車' in n or 'rac' in n or '競速' in n:
        return 'RAC 賽車'
    if 'spg' in n or 'spt' in n or '運動' in n or '棒球' in n or '高爾夫' in n or '足球' in n or '賽馬' in n or '籃球' in n:
        return 'SPT 運動'
    if 'puz' in n or '益智' in n or '方塊' in n:
        return 'PUZ 益智'
    if 'tbl' in n or '桌遊' in n or '麻將' in n or '麻雀' in n or '大富翁' in n or '柏青哥' in n or '小鋼珠' in n or '棋' in n:
        return 'TBL 桌遊'
    return 'ETC 其他'

def extract_series(jp, zh, en):
    t = f"{zh} {jp} {en}"
    if '瑪利歐' in t or 'Mario' in t: return '超級瑪利歐'
    if '勇者鬥惡龍' in t or 'Dragon Quest' in t: return '勇者鬥惡龍'
    if '太空戰士' in t or '最終幻想' in t or 'Final Fantasy' in t: return '最終幻想'
    if '聖劍傳說' in t or 'Seiken Densetsu' in t: return '聖劍傳說'
    if '復活邪神' in t or 'Romancing Sa-Ga' in t: return '復活邪神'
    if '七龍珠' in t or 'Dragon Ball' in t: return '七龍珠'
    if '亂馬' in t or 'Ranma' in t: return '亂馬1/2'
    if '熱血' in t or '國夫' in t or 'Kunio' in t: return '熱血高校'
    if '信長' in t or 'Nobunaga' in t: return '信長之野望'
    if '三國' in t or 'Sangokushi' in t: return '三國志'
    if '五右衛門' in t or 'Goemon' in t: return '大盜五右衛門'
    if '機戰' in t or '機器人大戰' in t or 'Super Robot' in t: return '超級機器人大戰'
    if '桃太郎' in t or 'Momotarou' in t: return '桃太郎'
    if '快打旋風' in t or 'Street Fighter' in t: return '快打旋風'
    if '洛克人' in t or 'Rockman' in t or 'Mega Man' in t: return '洛克人'
    if '德比賽馬' in t or 'Derby Stallion' in t: return '德比賽馬'
    if '賽馬大亨' in t or 'Winning Post' in t: return '賽馬大亨'
    if '薩爾達' in t or 'Zelda' in t: return '薩爾達傳說'
    if '聖火降魔錄' in t or 'Fire Emblem' in t: return '聖火降魔錄'
    if '超時空之鑰' in t or 'Chrono Trigger' in t: return '超時空之鑰'
    if '大戰略' in t or 'Daisenryaku' in t: return '大戰略'
    if '蒼狼與白鹿' in t or 'Aoki Ookami' in t: return '成吉思汗'
    if '魔法氣泡' in t or 'Puyo Puyo' in t: return '魔法氣泡'
    if '俄羅斯方塊' in t or 'Tetris' in t: return '俄羅斯方塊'
    if '巫術' in t or 'Wizardry' in t: return '巫術'
    return zh.split()[0] if zh else jp.split()[0]

def main():
    print("=== RetroVault SFC 全量管線建檔工具 ===")
    
    # Load SFC games from 收藏.xlsx
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, '收藏.xlsx'), data_only=True)
    ws = wb['SFC']

    sfc_games = []
    for r in range(2, ws.max_row+1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
        if any(vals) and vals[0] is not None:
            sfc_games.append({
                'no': vals[0],
                'jp': str(vals[1] or '').strip(),
                'zh': str(vals[2] or '').strip(),
                'en': str(vals[3] or '').strip(),
                'code': str(vals[4] or '').strip(),
                'region': str(vals[5] or '').strip(),
                'genre_raw': str(vals[6] or '').strip(),
            })

    print(f"[*] 載入 收藏.xlsx [SFC] 共 {len(sfc_games)} 款遊戲")

    # Match Libretro
    for g in sfc_games:
        b, s = match_libretro(g)
        g['libretro_box'] = b
        g['libretro_snap'] = s

    matched_b = sum(1 for g in sfc_games if g['libretro_box'])
    matched_s = sum(1 for g in sfc_games if g['libretro_snap'])
    print(f"[✓] Libretro 封面比對: {matched_b}/{len(sfc_games)}, 實機截圖比對: {matched_s}/{len(sfc_games)}")

    # Load / Init LaunchBox Cache
    cache_path = os.path.join(INTERNAL_DIR, 'sfc_lb_cache.json')
    lb_cache = {}
    if os.path.exists(cache_path):
        try:
            with open(cache_path, 'r', encoding='utf-8') as f:
                lb_cache = json.load(f)
            print(f"[*] 載入既有 LaunchBox 快取: {len(lb_cache)} 筆")
        except Exception:
            pass

    # LaunchBox Search function for a single game
    def fetch_lb_for_game(g):
        key = g['code'] or str(g['no'])
        if key in lb_cache and lb_cache[key].get('queried'):
            return key, lb_cache[key]

        b_q = query_from_boxart(g['libretro_box'])
        queries = []
        if b_q:
            queries.append(b_q)
            c_q = clean_lb_query(b_q)
            if c_q and c_q not in queries:
                queries.append(c_q)
        if g['en']:
            queries.append(clean_lb_query(g['en']))
        if g['jp']:
            queries.append(clean_lb_query(g['jp']))

        cands = []
        target_name = b_q if b_q else g['en']
        for q in queries:
            if not q: continue
            cands = search_launchbox(q)
            if cands:
                break

        if not cands:
            return key, {'queried': True, 'found': False}

        best = max(cands, key=lambda c: similarity(target_name, c['title']))
        imgs = get_launchbox_images(best['link'])
        res = {
            'queried': True,
            'found': True,
            'title': best['title'],
            'link': best['link'],
            'back': imgs.get('back'),
            'spine': imgs.get('spine'),
            'front': imgs.get('front'),
        }
        return key, res

    # Crawl LaunchBox with ThreadPoolExecutor
    to_crawl = [g for g in sfc_games if (g['code'] or str(g['no'])) not in lb_cache or not lb_cache[(g['code'] or str(g['no']))].get('queried')]
    if to_crawl:
        print(f"[*] 併發查詢 LaunchBox 圖資 (待查詢 {len(to_crawl)} 款)...")
        with ThreadPoolExecutor(max_workers=6) as executor:
            futures = {executor.submit(fetch_lb_for_game, g): g for g in to_crawl}
            cnt = 0
            for fut in as_completed(futures):
                k, res = fut.result()
                lb_cache[k] = res
                cnt += 1
                if cnt % 20 == 0 or cnt == len(to_crawl):
                    print(f"  已查詢 {cnt}/{len(to_crawl)} 款...")

        with open(cache_path, 'w', encoding='utf-8') as f:
            json.dump(lb_cache, f, ensure_ascii=False, indent=2)
        print(f"[✓] LaunchBox 查詢完成，快取寫入 {cache_path}")

    # Download back and spine images to docs/cdn/
    print("[*] 正在下載 back / spine 圖檔至 docs/cdn/ 並生成 Raw 網址...")
    for idx, g in enumerate(sfc_games):
        k = g['code'] or str(g['no'])
        info = lb_cache.get(k, {})
        back_src = info.get('back', {}).get('src') if info.get('back') else None
        spine_src = info.get('spine', {}).get('src') if info.get('spine') else None

        g['back_url'] = download_cdn_image(back_src) if back_src else ''
        g['spine_url'] = download_cdn_image(spine_src) if spine_src else ''

    total_backs = sum(1 for g in sfc_games if g['back_url'])
    total_spines = sum(1 for g in sfc_games if g['spine_url'])
    print(f"[✓] 本機 CDN 鏡像圖檔就緒: 封底 {total_backs}/{len(sfc_games)}, 書背 {total_spines}/{len(sfc_games)}")

    # 3. Construct RetroVault_遊戲_SFC.csv (54 columns)
    sfc_csv_path = os.path.join(DATA_DIR, 'RetroVault_遊戲_SFC.csv')
    raw_snes_base = "https://raw.githubusercontent.com/libretro-thumbnails/Nintendo_-_Super_Nintendo_Entertainment_System/master/"

    headers_en = [
        'uuid', 'category', 'primary_name', 'jp_name', 'zh_name', 'en_name', 'series', 'platform',
        'media', 'region', 'edition', 'voice_language', 'subtitle_language', 'age_rating', 'code',
        'barcode', 'developer', 'publisher', 'release_date', 'suggest_price', 'genre', 'players',
        'features', 'collect_status', 'completeness', 'loan_to', 'loan_due', 'bundle', 'bonus_code_status',
        'buy_date', 'buy_price', 'buy_source', 'storage_location', 'market_value', 'market_value_updated_at',
        'market_value_source', 'market_value_confidence', 'market_value_notes', 'market_value_currency',
        'market_value_foreign_amount', 'resale_price', 'resale_date', 'play_status', 'rating', 'summary',
        'ref_link', 'cover_img', 'back_img', 'spine_img', 'gameplay_img', 'extra_images', 'related_code',
        'notes', 'created_at'
    ]
    headers_zh = [
        'UUID', '資料類別', '主要名稱', '日文名稱', '中文名稱', '英文名稱', '系列名稱', '平台',
        '媒體類型', '遊戲區域', '版本說明', '語音語言', '字幕語言', '年齡分級', '遊戲編碼',
        '條碼 Barcode', '開發商', '發行商', '發行時間', '建議售價', '遊戲類型', '遊戲人數',
        '主要支援功能', '收藏狀態', '配件完整度', '借出對象', '預計歸還日', '特典/附件說明', 'DLC狀態',
        '購入日期', '購入價格', '入手來源', '保管位置', '目前估值', '估值更新日',
        '估值來源', '估值信心', '估值備註', '估值原幣', '原幣估值', '出售價格', '出售日期',
        '遊玩狀態', '個人評分', '遊戲簡介', '參考連結', '封面圖檔', '封底圖檔', '側邊圖檔',
        '', '自訂圖片', '關聯商品', '備註', '建檔時間'
    ]

    sfc_rows = []
    for g in sfc_games:
        u = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"sfc_{g['code']}_{g['no']}"))
        dev, pub = extract_developer_publisher(g['genre_raw'])
        genre = extract_genre(g['genre_raw'])
        series = extract_series(g['jp'], g['zh'], g['en'])

        cover_img = raw_snes_base + urllib.parse.quote(g['libretro_box']) if g['libretro_box'] else ''
        gameplay_img = raw_snes_base + urllib.parse.quote(g['libretro_snap']) if g['libretro_snap'] else ''
        back_img = g['back_url']
        spine_img = g['spine_url']
        ref_link = f"https://www.google.com/search?q={urllib.parse.quote(g['jp'] + ' Super Famicom')}"
        summary = f"Super Famicom (SFC) 經典作品《{g['zh']}》（日文原名：{g['jp']}），編號 {g['code']}。{g['genre_raw']}"

        row = [
            u,                              # uuid
            '遊戲',                         # category
            g['jp'],                        # primary_name
            g['jp'],                        # jp_name
            g['zh'],                        # zh_name
            g['en'],                        # en_name
            series,                         # series
            'Super Famicom',                # platform
            '卡匣（ROM 卡帶）',            # media
            '日本',                         # region
            '一般版',                       # edition
            '',                             # voice_language
            '',                             # subtitle_language
            '全區普遍級',                   # age_rating
            g['code'],                      # code
            '',                             # barcode
            dev,                            # developer
            pub,                            # publisher
            '1992/01/01',                   # release_date
            '0',                            # suggest_price
            genre,                          # genre
            '單人',                         # players
            '',                             # features
            '收藏中',                       # collect_status
            '外盒:良好|本體:良好|說明書:良好', # completeness
            '',                             # loan_to
            '',                             # loan_due
            '完整盒裝',                     # bundle
            'N/A',                          # bonus_code_status
            '2099/01/01',                   # buy_date
            '0',                            # buy_price
            '二手購入',                     # buy_source
            '',                             # storage_location
            '0',                            # market_value
            '',                             # market_value_updated_at
            '',                             # market_value_source
            'N/A',                          # market_value_confidence
            '',                             # market_value_notes
            '',                             # market_value_currency
            '',                             # market_value_foreign_amount
            '0',                            # resale_price
            '2099/01/01',                   # resale_date
            '未遊玩',                       # play_status
            'N/A',                          # rating
            summary,                        # summary
            ref_link,                       # ref_link
            cover_img,                      # cover_img
            back_img,                       # back_img
            spine_img,                      # spine_img
            gameplay_img,                   # gameplay_img
            '',                             # extra_images
            '',                             # related_code
            g['genre_raw'],                 # notes
            '2026-10-17 10:00:00'           # created_at
        ]
        sfc_rows.append(row)

    with open(sfc_csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(headers_en)
        writer.writerow(headers_zh)
        writer.writerows(sfc_rows)

    print(f"[✓] 已成功生成 {sfc_csv_path} (54 欄位, 包含 177 款 SFC 遊戲)")

    # 4. Merge FC (107) + MD (33) + SFC (177) into RetroVault_GS_Games_全量匯入檔.csv (57 columns)
    gs_csv_path = os.path.join(DATA_DIR, 'RetroVault_GS_Games_全量匯入檔.csv')
    gs_xlsx_path = os.path.join(DATA_DIR, 'RetroVault_GS_Games_全量匯入檔.xlsx')

    gs_headers = [
        'category', 'primary_name', 'jp_name', 'zh_name', 'en_name', 'series', 'platform', 'media',
        'region', 'edition', 'language', 'age_rating', 'code', 'barcode', 'developer', 'publisher',
        'release_date', 'suggest_price', 'genre', 'players', 'features', 'collect_status', 'completeness',
        'loan_to', 'loan_due', 'bundle', 'buy_date', 'buy_price', 'buy_source', 'play_status', 'rating',
        'summary', 'ref_link', 'cover_img', 'back_img', 'spine_img', 'gameplay_img', 'extra_images',
        'related_code', 'notes', 'uuid', 'created_at', 'storage_location', 'market_value', 'resale_price',
        'resale_date', 'dlc_status', 'bonus_code_status', 'voice_language', 'subtitle_language',
        'market_value_updated_at', 'market_value_source', 'market_value_confidence', 'market_value_notes',
        'market_value_currency', 'market_value_foreign_amount', 'favorite'
    ]

    all_gs_rows = []

    # Read FC (107)
    with open(os.path.join(DATA_DIR, 'RetroVault_遊戲_FC.csv'), 'r', encoding='utf-8') as f:
        r = csv.reader(f)
        next(r); next(r)
        for row in r:
            if not row or not row[0]: continue
            gs_row = [
                row[1],  # category
                row[2],  # primary_name
                row[3],  # jp_name
                row[4],  # zh_name
                row[5],  # en_name
                row[6],  # series
                row[7],  # platform
                row[8],  # media
                row[9],  # region
                row[10], # edition
                '',      # language
                row[13], # age_rating
                row[14], # code
                row[15], # barcode
                row[16], # developer
                row[17], # publisher
                row[18], # release_date
                row[19], # suggest_price
                row[20], # genre
                row[21], # players
                row[22], # features
                row[23], # collect_status
                row[24], # completeness
                row[25], # loan_to
                row[26], # loan_due
                row[27], # bundle
                row[29], # buy_date
                row[30], # buy_price
                row[31], # buy_source
                row[42], # play_status
                row[43], # rating
                row[44], # summary
                row[45], # ref_link
                row[46], # cover_img
                row[47], # back_img
                row[48], # spine_img
                row[49], # gameplay_img
                row[50], # extra_images
                row[51], # related_code
                row[52], # notes
                row[0],  # uuid
                row[53], # created_at
                row[32], # storage_location
                row[33], # market_value
                row[40], # resale_price
                row[41], # resale_date
                '',      # dlc_status
                row[28], # bonus_code_status
                row[11], # voice_language
                row[12], # subtitle_language
                row[34], # market_value_updated_at
                row[35], # market_value_source
                row[36], # market_value_confidence
                row[37], # market_value_notes
                row[38], # market_value_currency
                row[39], # market_value_foreign_amount
                ''       # favorite
            ]
            all_gs_rows.append(gs_row)

    # Read MD (33)
    with open(os.path.join(DATA_DIR, 'RetroVault_遊戲_MD.csv'), 'r', encoding='utf-8') as f:
        r = csv.reader(f)
        next(r); next(r)
        for row in r:
            if not row or not row[0]: continue
            gs_row = [
                row[1],  # category
                row[2],  # primary_name
                row[3],  # jp_name
                row[4],  # zh_name
                row[5],  # en_name
                row[6],  # series
                row[7],  # platform
                row[8],  # media
                row[9],  # region
                row[10], # edition
                '',      # language
                row[13], # age_rating
                row[14], # code
                row[15], # barcode
                row[16], # developer
                row[17], # publisher
                row[18], # release_date
                row[19], # suggest_price
                row[20], # genre
                row[21], # players
                row[22], # features
                row[23], # collect_status
                row[24], # completeness
                row[25], # loan_to
                row[26], # loan_due
                row[27], # bundle
                row[29], # buy_date
                row[30], # buy_price
                row[31], # buy_source
                row[42], # play_status
                row[43], # rating
                row[44], # summary
                row[45], # ref_link
                row[46], # cover_img
                row[47], # back_img
                row[48], # spine_img
                row[49], # gameplay_img
                row[50], # extra_images
                row[51], # related_code
                row[52], # notes
                row[0],  # uuid
                row[53], # created_at
                row[32], # storage_location
                row[33], # market_value
                row[40], # resale_price
                row[41], # resale_date
                '',      # dlc_status
                row[28], # bonus_code_status
                row[11], # voice_language
                row[12], # subtitle_language
                row[34], # market_value_updated_at
                row[35], # market_value_source
                row[36], # market_value_confidence
                row[37], # market_value_notes
                row[38], # market_value_currency
                row[39], # market_value_foreign_amount
                ''       # favorite
            ]
            all_gs_rows.append(gs_row)

    # Add SFC (177)
    for row in sfc_rows:
        gs_row = [
            row[1],  # category
            row[2],  # primary_name
            row[3],  # jp_name
            row[4],  # zh_name
            row[5],  # en_name
            row[6],  # series
            row[7],  # platform
            row[8],  # media
            row[9],  # region
            row[10], # edition
            '',      # language
            row[13], # age_rating
            row[14], # code
            row[15], # barcode
            row[16], # developer
            row[17], # publisher
            row[18], # release_date
            row[19], # suggest_price
            row[20], # genre
            row[21], # players
            row[22], # features
            row[23], # collect_status
            row[24], # completeness
            row[25], # loan_to
            row[26], # loan_due
            row[27], # bundle
            row[29], # buy_date
            row[30], # buy_price
            row[31], # buy_source
            row[42], # play_status
            row[43], # rating
            row[44], # summary
            row[45], # ref_link
            row[46], # cover_img
            row[47], # back_img
            row[48], # spine_img
            row[49], # gameplay_img
            row[50], # extra_images
            row[51], # related_code
            row[52], # notes
            row[0],  # uuid
            row[53], # created_at
            row[32], # storage_location
            row[33], # market_value
            row[40], # resale_price
            row[41], # resale_date
            '',      # dlc_status
            row[28], # bonus_code_status
            row[11], # voice_language
            row[12], # subtitle_language
            row[34], # market_value_updated_at
            row[35], # market_value_source
            row[36], # market_value_confidence
            row[37], # market_value_notes
            row[38], # market_value_currency
            row[39], # market_value_foreign_amount
            ''       # favorite
        ]
        all_gs_rows.append(gs_row)

    print(f"[*] 全量匯入檔整合完成: FC ({107}) + MD ({33}) + SFC ({len(sfc_rows)}) = {len(all_gs_rows)} 款遊戲")

    with open(gs_csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(gs_headers)
        writer.writerows(all_gs_rows)
    print(f"[✓] 寫入 {gs_csv_path} (57 欄位, {len(all_gs_rows)} 筆資料)")

    if openpyxl:
        wb_gs = openpyxl.Workbook()
        ws_gs = wb_gs.active
        ws_gs.title = 'RetroVault_Games'
        ws_gs.append(gs_headers)
        for r in all_gs_rows:
            ws_gs.append(r)
        wb_gs.save(gs_xlsx_path)
        print(f"[✓] 重新生成 {gs_xlsx_path}")

    print("=== SFC 全量管線處理完畢 ===")

if __name__ == '__main__':
    main()
